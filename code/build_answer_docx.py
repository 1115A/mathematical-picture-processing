from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = next(ROOT.glob("*作答版.docx"))


def main():
    doc = Document(str(DOCX_PATH))
    normal_style = doc.styles["Normal"]
    normal_font = normal_style.font.name or "Times New Roman"
    normal_size = normal_style.font.size

    def set_run_font(run):
        run.font.name = normal_font
        if normal_size:
            run.font.size = normal_size
        if run._element.rPr is not None:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), normal_font)

    def insert_paragraph_after(paragraph, text="", style="Normal", bold=False):
        new_p = OxmlElement("w:p")
        paragraph._p.addnext(new_p)
        p = Paragraph(new_p, paragraph._parent)
        p.style = style
        if text:
            run = p.add_run(text)
            run.bold = bold
            set_run_font(run)
        return p

    def insert_image_after(paragraph, image_path, caption):
        cursor = insert_paragraph_after(paragraph, caption)
        pic_p = insert_paragraph_after(cursor)
        run = pic_p.add_run()
        run.add_picture(str(image_path), width=Inches(5.8))
        return pic_p

    def find_anchor(text):
        for p in doc.paragraphs:
            if p.text.strip() == text:
                return p
        raise ValueError(f"Anchor not found: {text[:60]}")

    def code_text(name):
        return (ROOT / "code" / name).read_text(encoding="utf-8")

    def add_experiment_block(anchor_text, title, code_files, image_caps, analysis):
        cursor = find_anchor(anchor_text)
        cursor = insert_paragraph_after(cursor, f"{title} 实验结果与分析：", bold=True)
        for line in analysis:
            cursor = insert_paragraph_after(cursor, line)
        for img, cap in image_caps:
            cursor = insert_image_after(cursor, ROOT / "results" / img, cap)
        for fn in code_files:
            cursor = insert_paragraph_after(cursor, f"程序清单：{fn}", bold=True)
            cursor = insert_paragraph_after(cursor, code_text(fn))

    blocks = [
        (
            "（2）滤波模板大小对去噪效果的影响。",
            "4.滤波器去噪效果",
            ["exp3_4_custom_filters.m"],
            [
                ("exp3_4_custom_mean_filter.jpg", "图4-1 自编均值滤波器对高斯噪声和椒盐噪声的处理结果"),
                ("exp3_4_custom_median_filter.jpg", "图4-2 自编中值滤波器对高斯噪声和椒盐噪声的处理结果"),
            ],
            [
                "实验采用 electric.tif，分别加入方差为0.2的高斯噪声和概率为0.2的椒盐噪声，然后使用自编均值滤波器与自编中值滤波器进行3x3、5x5、9x9模板处理。",
                "定性结果表明：均值滤波对高斯噪声有较明显的平滑作用，但会削弱边缘和细节；中值滤波对椒盐噪声效果更好，能够有效去除黑白脉冲点，同时较好保持边缘。模板越大，去噪越强，但图像越容易模糊，9x9模板细节损失最明显。",
            ],
        ),
        (
            "（3）将原图像和两种滤波的图像在一个图形窗口显示，并比较两种滤波器滤波的效果。",
            "3.图像的锐化滤波",
            ["exp3_3_sharpening.m"],
            [("exp3_3_sharpening_and_mean.jpg", "图3-1 原图、Laplacian锐化结果和均值滤波结果对比")],
            [
                "Laplacian滤波属于二阶微分锐化处理，能够突出建筑边缘、线条和灰度突变区域，使轮廓更加清晰；均值滤波属于平滑处理，能够抑制局部灰度起伏，但会使边缘和纹理变模糊。"
            ],
        ),
        (
            "（3）利用imnoise()函数在图像electric.tif 上加入高斯噪声（均值为0，方差为0.2）和椒盐噪声（概率为0.2），利用medfilt2() 函数对加了高斯和椒盐噪声的图像分别进行均值滤波和中值滤波处理，模板为3✕3，比较中值滤波和均值滤波的效果。",
            "2.空间滤波去除噪声",
            ["exp3_2_spatial_denoising.m"],
            [
                ("exp3_2_median_same_template.jpg", "图2-1 3x3中值滤波对不同概率椒盐噪声的处理结果"),
                ("exp3_2_median_different_templates.jpg", "图2-2 不同模板大小的中值滤波结果"),
                ("exp3_2_mean_vs_median.jpg", "图2-3 均值滤波与中值滤波对高斯噪声和椒盐噪声的处理对比"),
            ],
            [
                "对椒盐噪声使用相同3x3模板时，噪声概率越大，残留噪声越多，图像细节越难恢复；对概率0.5的椒盐噪声，模板由3x3增大到9x9时，噪声点明显减少，但边缘和细节也逐渐模糊。",
                "对高斯噪声，均值滤波能够通过邻域平均降低随机波动，中值滤波也有一定平滑效果但不如均值滤波自然；对椒盐噪声，中值滤波明显优于均值滤波，因为中值运算能抑制极端黑白噪声点，而均值滤波会把脉冲噪声扩散成灰斑。",
            ],
        ),
        (
            "（2）利用imnoise()函数在图像electric.tif 上加入椒盐噪声，概率分别为0.2，0.5，0.8，在同一个图形窗口显示原图像和加入不同概率噪声的图像，比较加入不同概率的椒盐噪声对图像的影响。",
            "1.噪声生成",
            ["exp3_1_noise_generation.m"],
            [
                ("exp3_1_gaussian_noise.jpg", "图1-1 electric.tif加入不同方差高斯噪声的结果"),
                ("exp3_1_salt_pepper_noise.jpg", "图1-2 electric.tif加入不同概率椒盐噪声的结果"),
            ],
            [
                "高斯噪声表现为整幅图像灰度的随机扰动，方差从0.2增大到0.8时，图像颗粒感增强，局部细节和对比关系被随机波动淹没得更明显。",
                "椒盐噪声表现为随机黑白脉冲点，概率从0.2增大到0.8时，黑白噪声点密度快速增加，图像结构被大量极值像素破坏，主观可辨性明显下降。",
            ],
        ),
    ]

    for block in blocks:
        add_experiment_block(*block)

    answers = [
        (
            "1.图像灰度变换与空间滤波有什么不同?",
            "回答：图像灰度变换是点运算，输出像素主要由该位置输入像素的灰度值决定，常用于亮度、对比度、阈值和灰度级调整；空间滤波是邻域运算，输出像素由其周围模板内多个像素共同决定，可用于平滑去噪、边缘增强和锐化等处理。",
        ),
        (
            "2.空间锐化滤波和空间平滑滤波的特点和效果。",
            "回答：空间平滑滤波通常相当于低通处理，可削弱噪声和细小灰度起伏，使图像更平滑，但会带来边缘模糊和细节损失；空间锐化滤波通常突出高频成分和灰度突变，可增强边缘、轮廓和纹理，但也可能同步放大噪声。",
        ),
        (
            "3.结合实验内容，定性评价中值滤波和均值滤波对数字图像的高斯噪声和椒盐噪声的去除效果。",
            "回答：均值滤波对高斯噪声较有效，因为邻域平均能降低零均值随机扰动，但会模糊边缘；中值滤波对椒盐噪声更有效，因为中值可以排除极端黑白脉冲点并较好保留边缘。对于高斯噪声，中值滤波也能平滑但效果通常不如均值滤波自然；对于椒盐噪声，均值滤波会扩散噪声点，去噪效果不如中值滤波。",
        ),
        (
            "4.结合实验内容，定性评价滤波模板大小对去噪效果的影响？",
            "回答：滤波模板越大，参与计算的邻域像素越多，平滑和去噪能力越强，但图像细节、纹理和边缘也越容易被削弱。实验中3x3模板细节保持较好但对强噪声抑制有限，5x5在去噪和细节之间较折中，9x9去噪更明显但图像模糊最严重。实际应用中应根据噪声强度和细节保留要求选择模板大小。",
        ),
    ]

    for question, answer in answers:
        insert_paragraph_after(find_anchor(question), answer)

    doc.save(str(DOCX_PATH))
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
